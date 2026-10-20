"""
Planilha de indicações de rotas (plano de rotas) a partir dos dados do painel.
Uso: python planilha_plano.py <pasta_json_do_painel> <saida.xlsx>
     (a pasta com municipios.json, loc_<UF>.json e prop_<UF>.json gerada por preparar_painel.py)
"""
import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PASTA, OUT = sys.argv[1], sys.argv[2]
UFS = [("CE", "Ceará"), ("RN", "Rio Grande do Norte"), ("PB", "Paraíba"), ("PE", "Pernambuco"),
       ("AL", "Alagoas"), ("SE", "Sergipe"), ("BA", "Bahia"), ("PI", "Piauí"), ("MA", "Maranhão")]
ORD_UF = {u: i for i, (u, _) in enumerate(UFS)}
ler = lambda nome: json.load(open(os.path.join(PASTA, nome + ".json"), encoding="utf-8"))

M = ler("municipios")
mun = {m["chave"]: m for m in M}
P = [p for u, _ in UFS for p in ler("prop_" + u)]
L = [l for u, _ in UFS for l in ler("loc_" + u) if l["situacao"] != "Atendida"]

F = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
LINK = Font(name=F, size=10, color="0563C1", underline="single")
THIN = Border(bottom=Side(style="thin", color="D0D7E2"))
TOP = Border(top=Side(style="medium", color="1F3A5F"))
PFILL = {1: PatternFill("solid", fgColor="D5EFE9"), 2: PatternFill("solid", fgColor="DCE6F5"), 3: PatternFill("solid", fgColor="F5DCE8")}
INT, KM, H = "#,##0", "#,##0.0", "#,##0.0"

wb = Workbook()


def sheet(title, cols, rows, widths=None, prio_col=None):
    ws = wb.create_sheet(title)
    for j, (h, _, _) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, 2):
        for j, (h, get, fmt) in enumerate(cols, 1):
            v = get(r, i) if callable(get) and get.__code__.co_argcount == 2 else get(r)
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY
            if fmt == "link":
                if v:
                    c.hyperlink, c.font = v, LINK
            elif fmt:
                c.number_format = fmt
        if prio_col and r.get("prioridade"):
            ws.cell(row=i, column=prio_col).fill = PFILL[r["prioridade"]]
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    for j, (h, _, _) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(h, 14)
    return ws


# ---------- Rotas sugeridas ----------
P.sort(key=lambda p: (ORD_UF[p["uf"]], p["municipio"], p["prioridade"], p["km_por_localidade"] or 1e9))
g = lambda k: (lambda r: r.get(k))
pcols = [
    ("UF", g("uf"), None),
    ("Município", g("municipio"), None),
    ("Prioridade", g("prioridade"), None),
    ("Prioridade (descrição)", g("prioridade_txt"), None),
    ("Tipo", g("tipo"), None),
    ("Proposta", g("proposta"), None),
    ("Percurso sugerido", g("percurso"), None),
    ("Localidades atendidas", g("localidades_atendidas"), None),
    ("Nº de localidades", g("n_localidades"), INT),
    ("Km acrescentados", g("extensao_km"), KM),
    ("Km por localidade", lambda r, i: f'=IF(OR(I{i}=0,J{i}=""),"",J{i}/I{i})', KM),
    ("Tempo estimado", g("tempo"), None),
    ("Tempo (min)", g("tempo_min"), INT),
    ("Zona Eleitoral", lambda r: mun.get(r["uf"] + "|" + r["municipio"], {}).get("zona"), None),
    ("Juiz(a) Eleitoral", lambda r: mun.get(r["uf"] + "|" + r["municipio"], {}).get("juiz_eleitoral"), None),
    ("Observação para o pedido", lambda r: r.get("observacao") or None, None),
    ("Relatório PDF", lambda r: mun.get(r["uf"] + "|" + r["municipio"], {}).get("url_pdf"), "link"),
]
sheet("Rotas sugeridas", pcols, P, prio_col=3, widths={
    "UF": 6, "Município": 24, "Prioridade": 10, "Prioridade (descrição)": 22, "Tipo": 13, "Proposta": 34,
    "Percurso sugerido": 60, "Localidades atendidas": 40, "Juiz(a) Eleitoral": 30, "Observação para o pedido": 60,
    "Relatório PDF": 22})
NP = len(P) + 1
pc = {h: get_column_letter(i) for i, (h, _, _) in enumerate(pcols, 1)}
PR = lambda h: f"'Rotas sugeridas'!${pc[h]}$2:${pc[h]}${NP}"

# ---------- Localidades sem ônibus ----------
L.sort(key=lambda l: (ORD_UF[l["uf"]], l["municipio"], l["prioridade"] or 9, -(l["km_local_votacao"] or 0)))
lcols = [
    ("UF", g("uf"), None),
    ("Município", g("municipio"), None),
    ("Localidade", g("localidade"), None),
    ("Situação hoje", g("situacao"), None),
    ("Situação (como no relatório)", g("situacao_detalhe"), None),
    ("Km até o local de votação", g("km_local_votacao"), KM),
    ("Km até o ônibus mais próximo", g("km_ate_rota"), KM),
    ("Rota do edital mais próxima", g("rota_mais_proxima"), None),
    ("Indicação (rota sugerida)", lambda r: r.get("solucao_proposta") or "Sem proposta no relatório", None),
    ("Prioridade da indicação", g("prioridade"), None),
]
sheet("Localidades sem ônibus", lcols, L, prio_col=10, widths={
    "UF": 6, "Município": 24, "Localidade": 32, "Situação hoje": 16, "Situação (como no relatório)": 36,
    "Rota do edital mais próxima": 22, "Indicação (rota sugerida)": 44, "Prioridade da indicação": 12})
NL = len(L) + 1
lc = {h: get_column_letter(i) for i, (h, _, _) in enumerate(lcols, 1)}
LR = lambda h: f"'Localidades sem ônibus'!${lc[h]}$2:${lc[h]}${NL}"

# ---------- Por município ----------
MM = sorted([m for m in M if not m["fora_do_levantamento"] and (m["perto"] + m["sem_rota"]) > 0],
            key=lambda m: (ORD_UF[m["uf"]], m["nome"]))
crit = lambda i: f"$A{i},'Rotas sugeridas'!$A$2:$A${NP},$A{i},{PR('Município')},$B{i}"
mcols = [
    ("UF", g("uf"), None),
    ("Município", g("nome"), None),
    ("Zona Eleitoral", g("zona"), None),
    ("Juiz(a) Eleitoral", g("juiz_eleitoral"), None),
    ("Situação das rotas", g("situacao_rotas"), None),
    ("Localidades sem ônibus", lambda m, i: f"=COUNTIFS({LR('UF')},$A{i},{LR('Município')},$B{i})", INT),
    ("Prioridade 1 (ação rápida)", lambda m, i: f"=COUNTIFS({PR('UF')},$A{i},{PR('Município')},$B{i},{PR('Prioridade')},1)", INT),
    ("Prioridade 2 (rota nova eficiente)", lambda m, i: f"=COUNTIFS({PR('UF')},$A{i},{PR('Município')},$B{i},{PR('Prioridade')},2)", INT),
    ("Prioridade 3 (rota nova cara)", lambda m, i: f"=COUNTIFS({PR('UF')},$A{i},{PR('Município')},$B{i},{PR('Prioridade')},3)", INT),
    ("Km acrescentados", lambda m, i: f"=SUMIFS({PR('Km acrescentados')},{PR('UF')},$A{i},{PR('Município')},$B{i})", KM),
    ("Horas de viagem", lambda m, i: f"=SUMIFS({PR('Tempo (min)')},{PR('UF')},$A{i},{PR('Município')},$B{i})/60", H),
    ("Relatório PDF", g("url_pdf"), "link"),
]
sheet("Por município", mcols, MM, widths={"UF": 6, "Município": 24, "Juiz(a) Eleitoral": 30, "Situação das rotas": 32,
                                          "Relatório PDF": 22})
NMM = len(MM) + 1

# ---------- Resumo ----------
ws = wb["Sheet"]
ws.title = "Resumo"
ws.sheet_view.showGridLines = False
ws["A1"] = "Plano de rotas – indicações para localidades sem ônibus (Nordeste, 2º turno 2026)"
ws["A1"].font = Font(name=F, bold=True, size=14, color="1F3A5F")
ws["A2"] = ("Rotas sugeridas nos relatórios de fabio1974.github.io/eleicoes (19/10/2026), organizadas por prioridade. "
            "Confirme com o Cartório Eleitoral de cada zona antes de enviar o pedido.")
ws["A2"].font = Font(name=F, size=9, italic=True, color="555555")
rcols = [
    ("UF", None, None), ("Estado", None, None),
    ("Localidades sem ônibus", lambda u: f"=COUNTIFS({LR('UF')},{u})", INT),
    ("Cobertas por uma indicação", lambda u: f'=COUNTIFS({LR("UF")},{u},{LR("Prioridade da indicação")},"<>")', INT),
    ("Prioridade 1 (ação rápida)", lambda u: f"=COUNTIFS({PR('UF')},{u},{PR('Prioridade')},1)", INT),
    ("Prioridade 2 (rota nova eficiente)", lambda u: f"=COUNTIFS({PR('UF')},{u},{PR('Prioridade')},2)", INT),
    ("Prioridade 3 (rota nova cara)", lambda u: f"=COUNTIFS({PR('UF')},{u},{PR('Prioridade')},3)", INT),
    ("Total de indicações", "sum3", INT),
    ("Km acrescentados", lambda u: f"=SUMIFS({PR('Km acrescentados')},{PR('UF')},{u})", KM),
    ("Horas de viagem", lambda u: f"=SUMIFS({PR('Tempo (min)')},{PR('UF')},{u})/60", H),
]
HR = 4
for j, (h, _, _) in enumerate(rcols, 1):
    c = ws.cell(row=HR, column=j, value=h)
    c.font, c.fill = HEAD_FONT, HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if j > 2 else "left")
ws.row_dimensions[HR].height = 44
cl = {h: get_column_letter(j) for j, (h, _, _) in enumerate(rcols, 1)}
for i, (uf, nome) in enumerate(UFS):
    r = HR + 1 + i
    ws.cell(row=r, column=1, value=uf).font = BOLD
    ws.cell(row=r, column=2, value=nome).font = BODY
    for j, (h, f, fmt) in enumerate(rcols[2:], 3):
        v = (f"={cl['Prioridade 1 (ação rápida)']}{r}+{cl['Prioridade 2 (rota nova eficiente)']}{r}+{cl['Prioridade 3 (rota nova cara)']}{r}"
             if f == "sum3" else f(f"$A{r}"))
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.number_format, c.border = BODY, fmt, THIN
TR = HR + 1 + len(UFS)
ws.cell(row=TR, column=1, value="Total").font = BOLD
ws.cell(row=TR, column=2, value="Nordeste").font = BOLD
for j, (h, f, fmt) in enumerate(rcols[2:], 3):
    L_ = get_column_letter(j)
    c = ws.cell(row=TR, column=j, value=f"=SUM({L_}{HR + 1}:{L_}{TR - 1})")
    c.font, c.number_format, c.border = BOLD, fmt, TOP
ws.cell(row=TR, column=1).border = ws.cell(row=TR, column=2).border = TOP
ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 22
for j in range(3, len(rcols) + 1):
    ws.column_dimensions[get_column_letter(j)].width = 15

r = TR + 2
ws.cell(row=r, column=1, value="Critérios de prioridade").font = Font(name=F, bold=True, size=11, color="1F3A5F")
for k, txt in [
    (1, "1 – Ação rápida: parada extra ou extensão de uma rota do edital. Usa o ônibus que já roda; é o pedido mais simples."),
    (2, "2 – Rota nova eficiente: rota nova com até 10 km de percurso (ida e volta) por localidade atendida."),
    (3, "3 – Rota nova cara: rota nova com mais de 10 km por localidade atendida. Avaliar veículo e horário com o Cartório."),
]:
    r += 1
    c = ws.cell(row=r, column=1, value=k); c.font = BOLD; c.fill = PFILL[k]
    ws.cell(row=r, column=2, value=txt).font = BODY
r += 2
ws.cell(row=r, column=1, value="Como usar").font = Font(name=F, bold=True, size=11, color="1F3A5F")
for txt in [
    "Por município: quantas indicações de cada prioridade, km e horas – base para o ofício da Prefeitura ao Juiz Eleitoral.",
    "Rotas sugeridas: cada rota com percurso, localidades atendidas, km, tempo e a observação a incluir no pedido.",
    "Localidades sem ônibus: cada localidade que hoje fica sem transporte e a rota sugerida que a atende.",
    "A observação aparece quando o município teve o transporte dispensado, não tem rotas no DJE ou só tem rotas do 1º turno.",
    "Pela Lei nº 6.091/1974, só a Justiça Eleitoral define as rotas; a Prefeitura sugere percursos e cede veículos.",
]:
    r += 1
    ws.cell(row=r, column=2, value=txt).font = BODY

wb.move_sheet("Por município", offset=-(len(wb.sheetnames) - 2))
wb.save(OUT)
print("ok", len(P), len(L), len(MM))
