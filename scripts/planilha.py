"""
Monta a planilha a partir de dados.json (gerado por extrair.py).
Uso: python planilha.py <dados.json> <saida.xlsx>
"""
import sys
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open(sys.argv[1], encoding="utf-8"))
OUT = sys.argv[2]

UFS = [("CE", "Ceará"), ("RN", "Rio Grande do Norte"), ("PB", "Paraíba"), ("PE", "Pernambuco"),
       ("AL", "Alagoas"), ("SE", "Sergipe"), ("BA", "Bahia"), ("PI", "Piauí"), ("MA", "Maranhão")]

F = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
LINK = Font(name=F, size=10, color="0563C1", underline="single")
THIN = Border(bottom=Side(style="thin", color="D0D7E2"))
TOP = Border(top=Side(style="medium", color="1F3A5F"))

wb = Workbook()


def sheet(title, cols, rows, widths=None):
    ws = wb.create_sheet(title)
    for j, (h, _, fmt) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, 2):
        for j, (h, get, fmt) in enumerate(cols, 1):
            v = get(r)
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY
            if fmt == "link" and v:
                c.hyperlink = v
                c.font = LINK
            elif fmt and fmt != "link":
                c.number_format = fmt
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    for j, (h, get, _) in enumerate(cols, 1):
        w = (widths or {}).get(h)
        if w is None:
            sample = [len(str(get(r) or "")) for r in rows[:300]]
            w = min(max([len(h) / 2 + 2] + [s + 2 for s in sample]), 60)
        ws.column_dimensions[get_column_letter(j)].width = w
    return ws


pct = lambda v: None if v is None else v / 100
INT, P2, KM = "#,##0", "0.00%", "#,##0.0"
ordem_uf = {u: i for i, (u, _) in enumerate(UFS)}


def sit_resumida(s):
    if s.startswith("Atendida"):
        return "Atendida"
    if s.startswith("Rota passa perto"):
        return "Ônibus a 1–2 km"
    return "Sem rota"


# ---------- Municípios ----------
M = sorted(d["municipios"], key=lambda m: (ordem_uf[m["uf"]], m["nome"]))
g = lambda k: (lambda m: m.get(k))
mcols = [
    ("UF", g("uf"), None),
    ("Município", g("nome"), None),
    ("Estado", g("estado"), None),
    ("Zona Eleitoral", g("zona"), None),
    ("Juiz(a) Eleitoral", lambda m: m.get("juiz_eleitoral") or None, None),
    ("Situação das rotas", g("situacao_rotas"), None),
    ("Versão do relatório", lambda m: None if m.get("fora_do_levantamento") else m.get("versao"), None),
    ("Eleitores aptos", g("eleitores"), INT),
    ("Locais de votação", g("locais_votacao"), INT),
    ("Seções", g("secoes"), INT),
    ("Locais de votação rurais", g("locais_rurais"), INT),
    ("Eleitores em locais rurais", g("eleitores_rurais"), INT),
    ("% eleitores rurais", lambda m: pct(m.get("pct_eleitores_rurais")), "0%"),
    ("Localidades rurais analisadas", g("localidades_rurais_no_mapa"), INT),
    ("Localidades a mais de 5 km do local de votação", g("loc_mais_5km"), INT),
    ("Localidades a 3–5 km", g("loc_3a5km"), INT),
    ("Localidades que precisam de transporte", g("localidades_precisam"), INT),
    ("Rotas no edital", g("rotas_edital"), INT),
    ("Rotas novas propostas", g("rotas_novas"), INT),
    ("Paradas extras propostas", g("paradas_extras"), INT),
    ("Extensões de rota propostas", g("extensoes"), INT),
    ("Proposta (resumo)", g("proposta_resumo"), None),
    ("Abstenções 1º turno", g("abstencoes_1t"), INT),
    ("% abstenção 1º turno", lambda m: pct(m.get("abstencao_pct_1t")), P2),
    ("Votos Lula 1º turno", g("votos_Lula"), INT),
    ("% Lula", lambda m: pct(m.get("pct_Lula")), P2),
    ("Votos Flávio Bolsonaro 1º turno", g("votos_Flávio Bolsonaro"), INT),
    ("% Flávio Bolsonaro", lambda m: pct(m.get("pct_Flávio Bolsonaro")), P2),
    ("Edital de rotas", g("edital"), None),
    ("Horários do edital", g("horarios_edital"), None),
    ("Veículos do edital", g("veiculos_edital"), None),
    ("Observação sobre o turno", lambda m: m.get("observacao_turno") or None, None),
    ("Atualizado em", g("atualizado_em"), None),
    ("Página do município", g("url"), "link"),
    ("Relatório PDF", g("url_pdf"), "link"),
]
wsM = sheet("Municípios", mcols, M, widths={
    "UF": 6, "Município": 26, "Estado": 18, "Juiz(a) Eleitoral": 30, "Situação das rotas": 32,
    "Proposta (resumo)": 40, "Edital de rotas": 50, "Horários do edital": 45, "Veículos do edital": 35,
    "Observação sobre o turno": 50, "Página do município": 22, "Relatório PDF": 22})
for c in range(8, 29):
    wsM.column_dimensions[get_column_letter(c)].width = 13
NM = len(M) + 1

# ---------- Localidades ----------
L = sorted(d["localidades"], key=lambda l: (ordem_uf[l["uf"]], l["municipio"], -(l["km_ate_local_votacao"] or 0)))
lcols = [
    ("UF", lambda l: l["uf"], None),
    ("Município", lambda l: l["municipio"], None),
    ("Zona", lambda l: l["zona"], None),
    ("Localidade", lambda l: l["localidade"], None),
    ("Situação", lambda l: sit_resumida(l["situacao"]), None),
    ("Situação (como no relatório)", lambda l: l["situacao"], None),
    ("Distância ao local de votação (km, estrada)", lambda l: l["km_ate_local_votacao"], KM),
    ("Rota do edital mais próxima", lambda l: l["rota_mais_proxima"], None),
    ("Distância ao percurso da rota (km)", lambda l: l["km_ate_percurso_rota"], KM),
]
sheet("Localidades", lcols, L, widths={"UF": 6, "Município": 26, "Localidade": 34, "Situação": 18,
                                         "Situação (como no relatório)": 40, "Rota do edital mais próxima": 26})
NL = len(L) + 1

# ---------- Rotas propostas ----------
P = sorted(d["propostas"], key=lambda p: (ordem_uf[p["uf"]], p["municipio"], p["tipo_proposta"] != "Rota nova"))
pcols = [
    ("UF", lambda p: p["uf"], None),
    ("Município", lambda p: p["municipio"], None),
    ("Zona", lambda p: p["zona"], None),
    ("Tipo", lambda p: p["tipo_proposta"], None),
    ("Proposta", lambda p: p["proposta"], None),
    ("Percurso", lambda p: p["percurso"], None),
    ("Localidades atendidas", lambda p: p["localidades_atendidas"], None),
    ("Nº de localidades atendidas", lambda p: p["n_localidades_atendidas"], INT),
    ("Extensão (km)", lambda p: p["extensao_km"], KM),
    ("Extensão (como no relatório)", lambda p: p["extensao_texto"], None),
    ("Tempo estimado", lambda p: p["tempo"], None),
]
sheet("Rotas propostas", pcols, P, widths={"UF": 6, "Município": 26, "Proposta": 36, "Percurso": 70,
                                             "Localidades atendidas": 50, "Extensão (como no relatório)": 28})
NP = len(P) + 1

# ---------- Rotas do edital ----------
E = sorted(d["edital"], key=lambda e: (ordem_uf[e["uf"]], e["municipio"]))
sheet("Rotas do edital", [
    ("UF", lambda e: e["uf"], None),
    ("Município", lambda e: e["municipio"], None),
    ("Zona", lambda e: e["zona"], None),
    ("Edital", lambda e: e["edital"], None),
    ("Rota", lambda e: e["rota"], None),
    ("Percurso (na ordem do edital)", lambda e: e["percurso"], None),
], E, widths={"UF": 6, "Município": 26, "Edital": 50, "Rota": 14, "Percurso (na ordem do edital)": 100})

# ---------- Resumo por estado (primeira aba, com fórmulas) ----------
ws = wb["Sheet"]
ws.title = "Resumo"
ws.sheet_view.showGridLines = False
ws["A1"] = "Transporte de eleitores rurais – Nordeste (2º turno 2026)"
ws["A1"].font = Font(name=F, bold=True, size=14, color="1F3A5F")
n_rel = sum(1 for m in M if not m.get("fora_do_levantamento"))
ws["A2"] = (f"Dados extraídos dos {n_rel} relatórios municipais publicados em fabio1974.github.io/eleicoes "
            "(9 estados do Nordeste; relatórios atualizados em 19/10/2026).")
ws["A2"].font = Font(name=F, size=9, italic=True, color="555555")

mc = {h: get_column_letter(i) for i, (h, _, _) in enumerate(mcols, 1)}
lc = {h: get_column_letter(i) for i, (h, _, _) in enumerate(lcols, 1)}
pc = {h: get_column_letter(i) for i, (h, _, _) in enumerate(pcols, 1)}
MR = lambda h: f"'Municípios'!${mc[h]}$2:${mc[h]}${NM}"
LR = lambda h: f"Localidades!${lc[h]}$2:${lc[h]}${NL}"
PR = lambda h: f"'Rotas propostas'!${pc[h]}$2:${pc[h]}${NP}"

cols = [
    ("UF", None, None),
    ("Estado", None, None),
    ("Municípios com relatório", lambda u: f'=COUNTIFS({MR("UF")},{u},{MR("Versão do relatório")},"<>")', INT),
    ("Eleitores aptos", lambda u: f'=SUMIFS({MR("Eleitores aptos")},{MR("UF")},{u})', INT),
    ("Eleitores em locais rurais", lambda u: f'=SUMIFS({MR("Eleitores em locais rurais")},{MR("UF")},{u})', INT),
    ("Abstenções 1º turno", lambda u: f'=SUMIFS({MR("Abstenções 1º turno")},{MR("UF")},{u},{MR("Versão do relatório")},"<>")', INT),
    ("Localidades distantes (3 km+)", lambda u: f'=COUNTIFS({LR("UF")},{u})', INT),
    ("Atendidas", lambda u: f'=COUNTIFS({LR("UF")},{u},{LR("Situação")},"Atendida")', INT),
    ("Ônibus a 1–2 km", lambda u: f'=COUNTIFS({LR("UF")},{u},{LR("Situação")},"Ônibus a 1–2 km")', INT),
    ("Sem rota", lambda u: f'=COUNTIFS({LR("UF")},{u},{LR("Situação")},"Sem rota")', INT),
    ("% atendidas", "pct", "0.0%"),
    ("Rotas novas propostas", lambda u: f'=COUNTIFS({PR("UF")},{u},{PR("Tipo")},"Rota nova")', INT),
    ("Paradas extras propostas", lambda u: f'=COUNTIFS({PR("UF")},{u},{PR("Tipo")},"Parada extra")', INT),
    ("Km de rotas novas (ida e volta)", lambda u: f'=SUMIFS({PR("Extensão (km)")},{PR("UF")},{u},{PR("Tipo")},"Rota nova")', KM),
]
HR = 4
for j, (h, _, _) in enumerate(cols, 1):
    c = ws.cell(row=HR, column=j, value=h)
    c.font, c.fill = HEAD_FONT, HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if j > 2 else "left")
ws.row_dimensions[HR].height = 44
cl = {h: get_column_letter(j) for j, (h, _, _) in enumerate(cols, 1)}
for i, (uf, nome) in enumerate(UFS):
    r = HR + 1 + i
    ws.cell(row=r, column=1, value=uf).font = BOLD
    ws.cell(row=r, column=2, value=nome).font = BODY
    for j, (h, f, fmt) in enumerate(cols[2:], 3):
        if f == "pct":
            v = f'=IF({cl["Localidades distantes (3 km+)"]}{r}=0,0,{cl["Atendidas"]}{r}/{cl["Localidades distantes (3 km+)"]}{r})'
        else:
            v = f(f"$A{r}")
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.number_format, c.border = BODY, fmt, THIN
    ws.cell(row=r, column=1).border = ws.cell(row=r, column=2).border = THIN
TR = HR + 1 + len(UFS)
ws.cell(row=TR, column=1, value="Total").font = BOLD
ws.cell(row=TR, column=2, value="Nordeste").font = BOLD
for j, (h, f, fmt) in enumerate(cols[2:], 3):
    L_ = get_column_letter(j)
    if f == "pct":
        v = f'=IF({cl["Localidades distantes (3 km+)"]}{TR}=0,0,{cl["Atendidas"]}{TR}/{cl["Localidades distantes (3 km+)"]}{TR})'
    else:
        v = f"=SUM({L_}{HR + 1}:{L_}{TR - 1})"
    c = ws.cell(row=TR, column=j, value=v)
    c.font, c.number_format, c.border = BOLD, fmt, TOP
ws.cell(row=TR, column=1).border = ws.cell(row=TR, column=2).border = TOP
ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 22
for j in range(3, len(cols) + 1):
    ws.column_dimensions[get_column_letter(j)].width = 14
ws.freeze_panes = f"C{HR + 1}"

r = TR + 2
ws.cell(row=r, column=1, value="Como ler").font = Font(name=F, bold=True, size=11, color="1F3A5F")
for txt in [
    "Localidades distantes: comunidades rurais a 3 km ou mais, por estrada, do local de votação mais próximo.",
    "Atendidas: um ônibus do edital passa a até 1 km. Ônibus a 1–2 km: resolve-se com parada extra ou desvio.",
    "Sem rota: ônibus a mais de 2 km, município sem rotas no DJE ou transporte dispensado pela Justiça Eleitoral.",
    "Capitais ficam fora do levantamento (área essencialmente urbana); 6 municípios estavam com relatório em preparação.",
]:
    r += 1
    ws.cell(row=r, column=1, value=txt).font = BODY
r += 2
ws.cell(row=r, column=1, value="Abas").font = Font(name=F, bold=True, size=11, color="1F3A5F")
for nome, desc in [
    ("Municípios", "Uma linha por município: eleitorado, localidades, rotas, abstenção, juiz, edital e links."),
    ("Localidades", "Todas as localidades rurais distantes listadas nos relatórios, atendidas ou não."),
    ("Rotas propostas", "Cada rota nova, parada extra ou extensão proposta, com percurso, extensão e tempo."),
    ("Rotas do edital", "Rotas já publicadas pela Justiça Eleitoral, na ordem do edital."),
]:
    r += 1
    ws.cell(row=r, column=1, value=nome).font = BOLD
    ws.cell(row=r, column=3, value=desc).font = BODY
r += 2
nota = ("Fontes (segundo o site): TRE e TSE (locais de votação), Diário da Justiça Eletrônico (editais de rotas), "
        "IBGE – Localidades do Brasil 2022, OpenStreetMap e OSRM (distâncias por estrada). O próprio levantamento "
        "recomenda confirmar os números com o Cartório Eleitoral de cada zona.")
ws.cell(row=r, column=1, value=nota).font = Font(name=F, size=9, italic=True, color="555555")

wb.save(OUT)
print("ok")
